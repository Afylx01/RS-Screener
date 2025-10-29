import pandas as pd
from datetime import datetime
import os
import requests
from openpyxl.styles import PatternFill, Font, Alignment
from config import RESULTS_DIR, OUTPUT_SETTINGS, TELEGRAM_CONFIG, DELIVERY_SCANNER_DIR

def generate_tradingview_link(symbol):
    cleaned_symbol = symbol.replace(".NS", "")
    return f'https://www.tradingview.com/chart/?symbol=NSE:{cleaned_symbol}'

def format_results(results_df):
    if results_df.empty: return results_df
    for col in ['Close', 'EMA_50', 'EMA_200', 'RS55_Today', 'RS55_Yesterday', '52W_High_Ratio']:
        if col in results_df.columns:
            results_df[col] = results_df[col].round(2)
    if OUTPUT_SETTINGS['generate_tradingview_links']:
        results_df['TradingView_Link'] = results_df['Symbol'].apply(generate_tradingview_link)
    results_df['Symbol'] = results_df['Symbol'].str.replace(".NS", "", regex=False)
    return results_df

def save_results(results_df, scan_date):
    if results_df.empty:
        print(f"No stocks met the criteria for {scan_date.strftime('%Y-%m-%d')}.")
        return
    filename_date = scan_date.strftime('%Y%m%d')
    excel_path = RESULTS_DIR / f"RS55_Scan_{filename_date}.xlsx"
    results_df.to_excel(excel_path, index=False, sheet_name='Scan Results')
    print(f"Results saved to {excel_path}")

def print_top_results(results_df):
    if not results_df.empty:
        print("\n--- Top Results from Initial Scan ---")
        print(results_df.head(OUTPUT_SETTINGS['top_results_console']).to_string())

def save_advanced_scan_results(df_scan: pd.DataFrame, df_signals: pd.DataFrame, input_file: str):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base_name = os.path.splitext(os.path.basename(input_file))[0]
    filename = f"advanced_scan_{base_name}_{timestamp}.xlsx"
    output_path = "advanced_results"
    os.makedirs(output_path, exist_ok=True)

    with pd.ExcelWriter(os.path.join(output_path, filename), engine='openpyxl') as writer:
        if not df_scan.empty:
            df_scan.to_excel(writer, sheet_name='Full_Advanced_Scan', index=False)
        if not df_signals.empty:
            df_signals.to_excel(writer, sheet_name='Top_5_Signals', index=False)
            if 'shares' in df_signals.columns:
                portfolio_cols = [
                    'symbol', 'rank', 'score', 'hhhl_trend', 'trend_strength',
                    'entry_type', 'close', 'shares', 'position_value', 'stop_price',
                    'target_1r', 'target_2r', 'position_pct'
                ]
                available_cols = [col for col in portfolio_cols if col in df_signals.columns]
                df_signals[available_cols].to_excel(writer, sheet_name='Portfolio_Allocation', index=False)

        for sheet_name in writer.sheets:
            ws = writer.sheets[sheet_name]
            for cell in ws[1]:
                cell.fill = PatternFill(start_color='4CAF50', end_color='4CAF50', fill_type='solid')
                cell.font = Font(bold=True, color='FFFFFF')
                cell.alignment = Alignment(horizontal='center', vertical='center')
            for column in ws.columns:
                max_length = max(len(str(cell.value)) for cell in column)
                ws.column_dimensions[column[0].column_letter].width = min(max_length + 2, 50)
            ws.freeze_panes = 'A2'
    print(f"\nAdvanced scan results saved to: {os.path.join(output_path, filename)}")

def display_advanced_summary(df_scan: pd.DataFrame, df_signals: pd.DataFrame, portfolio_value: float):
    print("\n" + "="*60)
    print("📊 ADVANCED SCAN SUMMARY")
    print("="*60)
    if df_scan.empty:
        print("❌ No advanced scan results to display.")
        return

    print(f"\n📈 Total stocks analyzed: {len(df_scan)}")
    if 'hhhl_trend' in df_scan.columns:
        print(f"📈 Stocks in HHHL Uptrend: {(df_scan['hhhl_trend'] == 'UPTREND').sum()}")
    print(f"✅ Tradeable stocks (HHHL + Filters): {df_scan['tradeable'].sum()}")
    if 'entry_type' in df_scan.columns:
        print(f"🔥 Breakout signals: {(df_scan['entry_type'] == 'BREAKOUT').sum()}")
        print(f"📈 Pullback signals: {(df_scan['entry_type'] == 'PULLBACK').sum()}")

    if df_signals.empty:
        print("\n⚠️ No trading candidates found meeting all criteria.")
        return

    print("\n" + "="*60)
    print("🎯 TOP 5 TRADING CANDIDATES")
    print("="*60)
    for _, row in df_signals.iterrows():
        print(f"\n{row['rank']}. {row['symbol']}")
        print(f"   Score: {row['score']:.3f} | Price: ₹{row['close']:.2f}")
        print(f"   HHHL Trend: {row['trend']} (Strength: {row.get('trend_strength', 0):.1f}%)")
        print(f"   ADX: {row['adx']:.1f} | RSI: {row['rsi_14']:.1f} | RS55: {row['rs55']:.3f}")
        if row.get('entry_type'):
            print(f"   Signal: {row['entry_type']} 🎯")
            if 'shares' in row and row['shares'] > 0:
                print(f"   Position: {row['shares']} shares @ ₹{row['position_value']:,.0f} ({row['position_pct']}%)")
                print(f"   Stop: ₹{row['stop_price']:.2f} | T1: ₹{row['target_1r']:.2f} | T2: ₹{row['target_2r']:.2f}")

    if not df_signals.empty and 'position_value' in df_signals.columns:
        total_allocation = df_signals['position_value'].sum()
        print("\n" + "="*60)
        print("💰 PORTFOLIO ALLOCATION")
        print("="*60)
        print(f"Portfolio Value: ₹{portfolio_value:,.0f}")
        print(f"Total Allocation: ₹{total_allocation:,.0f} ({(total_allocation/portfolio_value)*100:.1f}%)")

def save_delivery_scan_results(df: pd.DataFrame, target_date: datetime):
    """Saves the high-delivery scan results to a formatted Excel file."""
    if df.empty:
        print("No stocks met the delivery scan criteria.")
        return None

    filename = f"High_Delivery_Scan_{target_date.strftime('%Y%m%d')}.xlsx"
    output_path = DELIVERY_SCANNER_DIR / filename

    # Add TradingView links
    df['tradingview_link'] = df['symbol'].apply(generate_tradingview_link)

    # Round numeric columns for cleaner output
    numeric_cols = [
        'close', 'prev_close', 'delivery_qty', 'pct_change',
        'avg_delivery_5d_prior', 'delivery_times', 'delivery_value_cr',
        'market_cap_cr', 'rs', 'rsi'
    ]
    for col in numeric_cols:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors='coerce').round(2)

    try:
        with pd.ExcelWriter(output_path, engine='xlsxwriter') as writer:
            df.to_excel(writer, sheet_name='High_Delivery_Stocks', index=False)
            # Apply formatting
            workbook = writer.book
            worksheet = writer.sheets['High_Delivery_Stocks']
            header_format = workbook.add_format({
                'bold': True, 'text_wrap': True, 'valign': 'top',
                'fg_color': '#4F81BD', 'font_color': 'white', 'border': 1
            })
            for col_num, value in enumerate(df.columns.values):
                worksheet.write(0, col_num, value, header_format)
            worksheet.autofit()
        print(f"Delivery scan results saved to: {output_path}")
        return output_path
    except Exception as e:
        print(f"Failed to save delivery scan results to Excel: {e}")
        return None

def send_telegram_report(file_path: str, caption: str):
    """Sends the specified file to the Telegram channel."""
    if not all([TELEGRAM_CONFIG['bot_token'], TELEGRAM_CONFIG['chat_id']]):
        print("Telegram credentials are not configured. Skipping notification.")
        return

    bot_token = TELEGRAM_CONFIG['bot_token']
    chat_id = TELEGRAM_CONFIG['chat_id']
    url = f"https://api.telegram.org/bot{bot_token}/sendDocument"

    try:
        with open(file_path, 'rb') as doc:
            files = {'document': doc}
            data = {'chat_id': chat_id, 'caption': caption}
            response = requests.post(url, data=data, files=files, timeout=20)

        if response.status_code == 200:
            print("Successfully sent report to Telegram.")
        else:
            print(f"Failed to send report to Telegram. Status: {response.status_code}, Response: {response.text}")
    except Exception as e:
        print(f"An error occurred while sending Telegram notification: {e}")
